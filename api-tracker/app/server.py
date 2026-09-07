"""FastAPI сервер для API Tracker.

Предоставляет:
- REST API для приёма логов вызовов (/api/log)
- Выдачу статистики и графиков (/api/stats)
- Поиск и фильтрацию логов (/api/logs)
- Экспорт в CSV и Excel (/api/export/excel, /api/export/csv)
- Раздачу статики дашборда (/)
"""

from __future__ import annotations

import csv
import io
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, Query, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from app.db import get_logs, get_stats, init_db, insert_call
from app.pricing import DEFAULT_PRICING, calculate_cost

STATIC_DIR = Path(__file__).resolve().parent / "static"
MSK_TZ = timezone(timedelta(hours=3))


def to_msk_str(ts_str: str) -> str:
    """Конвертировать ISO-таймстемп в московское время (MSK UTC+3)."""
    if not ts_str:
        return ""
    try:
        clean = ts_str.replace("Z", "+00:00")
        dt = datetime.fromisoformat(clean)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(MSK_TZ).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ts_str.replace("T", " ")[:19]


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    yield


app = FastAPI(title="API Tracker & Cost Monitor", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class LogEventRequest(BaseModel):
    provider: str = Field(..., description="vibecode | kie | google | openai | outsee | elevenlabs")
    model: str = Field(..., description="ID модели, например gemini-3.7-flash")
    call_type: str = Field("text", description="text | video | image | audio")
    key_alias: str = Field("", description="Маска или имя ключа")
    prompt_tokens: int = Field(0)
    completion_tokens: int = Field(0)
    cached_tokens: int = Field(0)
    media_count: int = Field(0)
    duration_sec: float = Field(0.0)
    cost_usd: float | None = Field(None, description="Если не передано, рассчитывается автоматически")
    status_code: int = Field(200)
    error_message: str = Field("")
    project_source: str = Field("", description="chat | pipeline | script")
    metadata: dict[str, Any] = Field(default_factory=dict)
    timestamp: str | None = Field(None)


@app.post("/api/log")
async def log_call(event: LogEventRequest):
    """Записать вызов API и рассчитать стоимость."""
    cost = event.cost_usd
    if cost is None:
        cost = calculate_cost(
            event.model,
            call_type=event.call_type,
            prompt_tokens=event.prompt_tokens,
            completion_tokens=event.completion_tokens,
            cached_tokens=event.cached_tokens,
            media_count=event.media_count,
            duration_sec=event.duration_sec,
            chars=int(event.metadata.get("chars", 0)),
        )

    call_id = insert_call(
        provider=event.provider,
        model=event.model,
        key_alias=event.key_alias,
        call_type=event.call_type,
        prompt_tokens=event.prompt_tokens,
        completion_tokens=event.completion_tokens,
        cached_tokens=event.cached_tokens,
        media_count=event.media_count,
        duration_sec=event.duration_sec,
        cost_usd=cost,
        status_code=event.status_code,
        error_message=event.error_message,
        project_source=event.project_source,
        metadata=event.metadata,
        timestamp=event.timestamp,
    )
    return {"status": "ok", "id": call_id, "cost_usd": cost}


@app.get("/api/stats")
async def stats(
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Сводные метрики и данные для графиков."""
    return get_stats(date_from=date_from, date_to=date_to)


@app.get("/api/logs")
async def logs_list(
    limit: int = Query(50, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    provider: str | None = Query(None),
    model: str | None = Query(None),
    call_type: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    search: str | None = Query(None),
):
    """Список логов активности с фильтрацией."""
    items, total = get_logs(
        limit=limit,
        offset=offset,
        provider=provider,
        model=model,
        call_type=call_type,
        status=status,
        date_from=date_from,
        date_to=date_to,
        search=search,
    )
    return {"items": items, "total": total, "limit": limit, "offset": offset}


@app.get("/api/pricing")
async def pricing_catalog():
    """Справочник цен по моделям."""
    return DEFAULT_PRICING


@app.get("/api/export/csv")
async def export_csv(
    provider: str | None = Query(None),
    model: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Экспорт логов в CSV."""
    items, _ = get_logs(
        limit=10000,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )

    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([
        "ID", "Время (МСК)", "Провайдер", "Модель", "Тип",
        "Входные токены", "Выходные токены", "Всего токенов",
        "Длительность (сек)", "Стоимость ($)", "Статус", "Источник", "Ошибка"
    ])

    for row in items:
        writer.writerow([
            row["id"],
            to_msk_str(row["timestamp"]),
            row["provider"],
            row["model"],
            row["call_type"],
            row["prompt_tokens"],
            row["completion_tokens"],
            row["total_tokens"],
            row["duration_sec"],
            f"{row['cost_usd']:.6f}",
            row["status_code"],
            row["project_source"],
            row["error_message"],
        ])

    data = buf.getvalue().encode("utf-8-sig")
    return Response(
        content=data,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=api_activity.csv"},
    )


@app.get("/api/export/excel")
async def export_excel(
    provider: str | None = Query(None),
    model: str | None = Query(None),
    status: str | None = Query(None),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
):
    """Экспорт логов в Excel (.xlsx) с форматированием."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    items, _ = get_logs(
        limit=10000,
        provider=provider,
        model=model,
        status=status,
        date_from=date_from,
        date_to=date_to,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Активность API"

    headers = [
        "ID", "Дата и время (МСК)", "Провайдер", "Модель", "Тип",
        "Входные токены", "Выходные токены", "Всего токенов",
        "Длительность (сек)", "Стоимость ($)", "Статус", "Источник", "Ошибка"
    ]
    ws.append(headers)

    # Оформление шапки
    header_fill = PatternFill(start_color="164E63", end_color="164E63", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in items:
        ws.append([
            row["id"],
            to_msk_str(row["timestamp"]),
            row["provider"],
            row["model"],
            row["call_type"],
            row["prompt_tokens"],
            row["completion_tokens"],
            row["total_tokens"],
            row["duration_sec"],
            round(row["cost_usd"], 6),
            row["status_code"],
            row["project_source"],
            row["error_message"],
        ])

    # Автоподбор ширины колонок
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=api_activity.xlsx"},
    )


# Главная страница дашборда
@app.get("/")
async def index():
    html_path = STATIC_DIR / "index.html"
    if html_path.is_file():
        return FileResponse(html_path)
    return {"message": "API Tracker Backend Running"}


if STATIC_DIR.is_dir():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
